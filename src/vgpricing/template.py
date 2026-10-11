"""Gera o template Excel de entrada (uma aba por tabela + LEIA-ME).

O Excel é só porta de entrada: nenhuma conta é feita na planilha.
Cabeçalhos = nomes da fonte definidos em config/colunas.yaml.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .schema import ESQUEMA

EXEMPLO = {
    "estipulantes": {"id_estipulante": "EST0001", "cnae_secao": "C", "cnae_divisao": "25", "data_inicio_vigencia": "01/07/2023",
                     "data_fim_vigencia": "", "tipo_capital": "multiplo_salarial", "contributario": "N",
                     "taxa_vigente_mensal_mil": 0.35},
    "vidas": {"id_estipulante": "EST0001", "id_vida": "a3f9c2e17b4d0e55", "data_nascimento": "15/06/1985",
              "sexo": "M", "data_inicio": "01/07/2023", "data_fim": "", "capital": 84000, "situacao": "ativo"},
    "sinistros": {"id_sinistro": "SIN000123", "id_estipulante": "EST0001", "id_vida": "a3f9c2e17b4d0e55",
                  "cobertura": "MQC", "data_ocorrencia": "10/03/2025", "data_aviso": "02/05/2025",
                  "data_pagamento": "20/05/2025", "valor_indenizado": 84000, "valor_reserva": 0, "status": "pago"},
    "premios": {"id_estipulante": "EST0001", "inicio_periodo": "01/07/2024", "fim_periodo": "30/06/2025",
                "premio_emitido": 152340.55, "taxa_mensal_mil": 0.35},
}

AJUDA = {
    "data_fim": "Vazio = vida em vigor. Último dia coberto (inclusive). Mudança de capital = nova linha.",
    "id_vida": "Identificador pseudonimizado (hash do CPF). Nunca o CPF em claro.",
    "capital": "Capital segurado de morte, em R$.",
    "situacao": "ativo | afastado | aposentado_invalidez",
    "sexo": "M ou F",
    "status": "pago | pendente | negado",
    "cnae_secao": "Seção da CNAE 2.0 (letra A a U) da atividade principal do estipulante.",
    "cnae_divisao": "Divisão da CNAE 2.0 (2 dígitos). Se vazia, usa-se o fator da seção.",
    "taxa_mensal_mil": "Taxa mensal por mil de capital vigente no período.",
    "valor_reserva": "Reserva de sinistro a liquidar (pendentes).",
}

FONTE = "Arial"


def criar_template_excel(caminho: str | Path, mapeamentos: dict) -> Path:
    wb = Workbook()
    leia = wb.active
    leia.title = "LEIA-ME"
    texto = [
        "Template de entrada - Vida em Grupo",
        "",
        "Como usar:",
        "1. Preencha as abas estipulantes e vidas (obrigatórias). sinistros e premios são opcionais na cotação de grupo novo.",
        "2. Cabeçalhos em amarelo são obrigatórios; em cinza, opcionais. Não renomeie cabeçalhos.",
        "3. A linha 2 de cada aba é um EXEMPLO: apague-a antes de importar.",
        "4. Datas em dd/mm/aaaa. Valores em R$ sem símbolo.",
        "5. Nenhuma conta é feita aqui: o cálculo roda em Python (vgpricing).",
        "6. Não inclua CPF, nome ou dados de saúde em claro (LGPD).",
    ]
    for i, t in enumerate(texto, 1):
        c = leia.cell(row=i, column=1, value=t)
        c.font = Font(name=FONTE, bold=(i == 1), size=12 if i == 1 else 10)
    leia.column_dimensions["A"].width = 110

    amarelo = PatternFill("solid", fgColor="FFFF00")
    cinza = PatternFill("solid", fgColor="D9D9D9")
    for tabela, esq in ESQUEMA.items():
        ws = wb.create_sheet(tabela)
        mapa = mapeamentos[tabela]
        cols = list(esq["obrigatorias"]) + list(esq["opcionais"])
        for j, canon in enumerate(cols, 1):
            cab = ws.cell(row=1, column=j, value=mapa.get(canon, canon))
            cab.font = Font(name=FONTE, bold=True)
            cab.fill = amarelo if canon in esq["obrigatorias"] else cinza
            cab.alignment = Alignment(horizontal="center")
            if canon in AJUDA:
                cab.comment = Comment(AJUDA[canon], "vgpricing")
            ex = ws.cell(row=2, column=j, value=EXEMPLO[tabela].get(canon, ""))
            ex.font = Font(name=FONTE, italic=True, color="808080")
            ws.column_dimensions[cab.column_letter].width = max(14, len(str(cab.value)) + 4)
        listas = {"sexo": '"M,F"', "situacao": '"ativo,afastado,aposentado_invalidez"',
                  "status": '"pago,pendente,negado"'}
        for canon, formula in listas.items():
            if canon in cols:
                letra = ws.cell(row=1, column=cols.index(canon) + 1).column_letter
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letra}2:{letra}1048576")
        ws.freeze_panes = "A2"
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
